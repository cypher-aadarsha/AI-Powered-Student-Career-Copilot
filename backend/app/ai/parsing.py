"""Extracts plain text from an uploaded resume file. Two formats are
supported — anything else is rejected by resume_service before it reaches
here (see resume_service.ALLOWED_MIME_TYPES).
"""
import io

from docx import Document
import pdfplumber

from app.core.exceptions import AppError

_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class UnsupportedResumeFormat(AppError):
    def __init__(self, message: str = "Unsupported resume file format."):
        super().__init__(message, status_code=422, code="unsupported_format")


def extract_text(file_bytes: bytes, mime_type: str) -> str:
    if mime_type == "application/pdf":
        return _extract_pdf_text(file_bytes)
    if mime_type == _DOCX_MIME:
        return _extract_docx_text(file_bytes)
    raise UnsupportedResumeFormat()


def _extract_pdf_text(file_bytes: bytes) -> str:
    text_parts: list[str] = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            text_parts.append(page.extract_text() or "")
    return "\n".join(text_parts).strip()


def _extract_docx_text(file_bytes: bytes) -> str:
    document = Document(io.BytesIO(file_bytes))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(p for p in parts if p).strip()
